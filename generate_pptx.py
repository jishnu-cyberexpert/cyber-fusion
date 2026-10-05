import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    # Color Palette: Cyber Dark Theme
    COLOR_BG = RGBColor(6, 9, 16)         # #060910
    COLOR_CARD = RGBColor(12, 16, 29)     # #0c101d
    COLOR_CYAN = RGBColor(6, 182, 212)    # #06b6d4
    COLOR_GREEN = RGBColor(16, 185, 129)  # #10b981
    COLOR_WHITE = RGBColor(248, 250, 252) # #f8fafc
    COLOR_MUTED = RGBColor(148, 163, 184) # #94a3b8
    COLOR_BORDER = RGBColor(26, 35, 58)   # #1a233a
    COLOR_AMBER = RGBColor(245, 158, 11)  # #f59e0b

    slides_data = [
        {
            "title": "CYBERFUSION XDR ENTERPRISE",
            "subtitle": "Unified Enterprise Cyber Defense & Security Operations Platform",
            "category": "PLATFORM OVERVIEW",
            "points": [
                "One Security Data Plane: Native integration of EDR, XDR, NDR, SIEM, SOAR, UEBA, TIP, and Digital Forensics.",
                "Zero Mockup & Real Telemetry: Operating entirely on real infrastructure streams, host events, and network packets.",
                "High-Throughput Streaming Bus: Sub-2ms processing latency with backpressure and Dead Letter Queue (DLQ).",
                "Strict Operating Modes: Complete segregation between LIVE production, LAB test harness, and DEMO environments."
            ]
        },
        {
            "title": "The Broken Security Operations Paradigm",
            "subtitle": "Critical Industry Challenges in Modern Enterprise SecOps",
            "category": "PROBLEM STATEMENT",
            "points": [
                "The 'Swivel-Chair' SOC: Analysts forced to manually toggle across 12+ isolated dashboards without shared context.",
                "Alert Fatigue & Correlation Failure: 10,000+ daily raw alarms; systems fail to link host activity with network C2 beacons.",
                "The 'Fake Security' Problem: Products claiming 'Attack Blocked' without real API integration at enforcement points.",
                "High Dwell Time: Batch database queries taking minutes/hours, granting attackers ample time to exfiltrate data."
            ]
        },
        {
            "title": "The Solution: One Security Data Plane",
            "subtitle": "Shared Real-Time Telemetry & Instant Normalization",
            "category": "ARCHITECTURE",
            "points": [
                "Unified Ingestion Bus: Telemetry streams from hosts, firewalls, and cloud APIs into a shared data plane.",
                "Standardized Schema: Automatic normalization into Open Cybersecurity Schema Framework (OCSF / ECS).",
                "Sub-millisecond TIP Enrichment: Immediate indicator matching against verified APT29, Lazarus, and LockBit CTI feeds.",
                "Measured Performance: 100,000+ EPS horizontal scalability with continuous queue depth and latency tracking."
            ]
        },
        {
            "title": "Real-Time Endpoint Detection & Response (EDR)",
            "subtitle": "Live Cross-Platform Host Agent Architecture",
            "category": "EDR CAPABILITIES",
            "points": [
                "Real Host Monitoring: Live Python daemon (psutil, socket, platform) inspecting active processes and network sockets.",
                "Full Process Lineage: Captures process creation, PID, PPID, parent process name, and full command line arguments.",
                "Mutual Auth & Fleet Management: Secure API token registration, heartbeat tracking, and operational health.",
                "Real Host Verified: Running live on host 'JISHNUPRASAD' (Windows 11, IP 192.168.1.37) with instant network isolation."
            ]
        },
        {
            "title": "Multi-Engine Production Detection Framework",
            "subtitle": "Sigma Rules, IOC Matching & Sliding Window Bursts",
            "category": "DETECTION ENGINE",
            "points": [
                "Sigma-Compatible Rules: Native execution of process rules (PowerShell cradles, Mimikatz LSASS access).",
                "Stateful Threshold Windows: Dynamic burst detection (e.g. IAM-T1110-001: 5+ failed authentications in 60s).",
                "IOC Network Detections: Outbound C2 beaconing alerts mapped to MITRE ATT&CK T1071.001.",
                "MITRE ATT&CK Alignment: Every detection mapped to Tactic, Technique, Sub-technique with false-positive guidance."
            ]
        },
        {
            "title": "Cross-Domain XDR Correlation Engine",
            "subtitle": "Entity Relationship Graph Matching Across Domains",
            "category": "XDR CORRELATION",
            "points": [
                "Graph Intersection: Links multi-domain signals based on shared Entity Graph (Hosts, Users, IPs, Hashes).",
                "Not Timestamp-Only: Prevents false correlations by requiring common entity attribution within a 2-hour window.",
                "Multi-Stage Killchain Consolidation: Merges Phishing -> Suspicious Process -> C2 Beacon -> Cloud Login into 1 Case.",
                "Reduced Ticket Burden: Consolidates hundreds of raw alerts into single actionable enterprise incidents."
            ]
        },
        {
            "title": "Contextual & Explainable Risk Engine",
            "subtitle": "Transparent Numeric Scoring (0 to 100)",
            "category": "RISK ANALYTICS",
            "points": [
                "Explainable Attribution: Eliminates black-box ML guesswork by providing transparent factor breakdowns.",
                "Severity Base Rating: Critical (+50), High (+35), Medium (+20).",
                "Privilege Amplification: +20 points when privileged identity (Domain Admin, root, svc_account) is targeted.",
                "Killchain Impact: +15 points for Command and Control, Exfiltration, or Impact tactics."
            ]
        },
        {
            "title": "User & Entity Behavior Analytics (UEBA)",
            "subtitle": "Welford's Method Baseline Profiling & Impossible Travel",
            "category": "BEHAVIORAL AI",
            "points": [
                "Dynamic Statistical Baselines: Numerically stable running mean and standard deviation per identity.",
                "Unusual Logon Hours: Detects authentications occurring outside normal working distribution (Z-score > 2.0).",
                "Impossible Travel Velocity: Mathematical delta calculation flagging logins from disparate IPs (>1,200 km/h).",
                "Transparent Cards: Shows observed value, historical mean, standard deviation, and calculated risk score."
            ]
        },
        {
            "title": "SOAR Response & 'No False Claims' Policy",
            "subtitle": "Response Orchestration with Strict Integration Verification",
            "category": "AUTOMATION & SOAR",
            "points": [
                "Automated & Analyst Playbooks: Instant host containment, process termination, and perimeter IP blocking.",
                "Human Authorization Gates: Mandatory Tier-2/Tier-3 approval triggers for high-impact actions.",
                "Strict Requirement: NO FALSE CLAIMS: Never claims 'IP Blocked' or 'Protected' if integration is missing.",
                "Real Enforcement Output: Explicitly reports 'ACTION NOT EXECUTED — INTEGRATION NOT CONFIGURED' with guide."
            ]
        },
        {
            "title": "Attack Surface Management & CNAPP",
            "subtitle": "Authorized Perimeter Discovery & Cloud Posture",
            "category": "SURFACE & CLOUD",
            "points": [
                "Authorized Scope Gate: Mandatory policy confirmation prior to external asset scanning.",
                "Automated Exposure Discovery: Probes open ports (80, 443, 22, 3389, 8080) and inspects SSL certificate validity.",
                "Multi-Cloud Posture (CNAPP): Live posture audit across AWS CloudTrail/S3, Azure NSG/Defender, GCP IAM.",
                "Remediation Guidance: Provides copy-paste AWS CLI, Azure CLI, and gcloud commands for instant remediation."
            ]
        },
        {
            "title": "Cryptographic Tamper-Resistant Audit Trail",
            "subtitle": "SHA-256 Blockchain-Style Chained Verification",
            "category": "COMPLIANCE & AUDIT",
            "points": [
                "Immutable Chained Blocks: Current Hash = SHA256(Previous Hash + Serialized Action Payload).",
                "Complete Operational Coverage: Logs every containment action, status change, and approval gate.",
                "One-Click Chain Verifier: Server recomputes hashes from Genesis Root; confirms zero database tampering.",
                "Regulatory Compliance: Technical evidence satisfying ISO 27001, SOC 2 Type II, and PCI DSS Requirement 10."
            ]
        },
        {
            "title": "Live Verification & Operational Highlights",
            "subtitle": "Real Fleet Demonstration & Active Performance",
            "category": "LIVE VERIFICATION",
            "points": [
                "Live Host Enrolled: Real Windows 11 host 'JISHNUPRASAD' actively streaming live processes and network sockets.",
                "Live WebSockets: Telemetry data plane and alerts feed updating dynamically without page refreshes.",
                "Sub-2ms Latency: End-to-end ingestion, normalization, detection, and correlation measured in real time.",
                "Commercial Console: Dark glassmorphic SOC console with 16 module views ready for enterprise operations."
            ]
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # Background fill
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.fill.background()

        # Top Accent Header Bar
        accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(0.06))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = COLOR_CYAN
        accent_bar.line.fill.background()

        # Category Pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(5), Inches(0.4))
        cat_tf = cat_box.text_frame
        p_cat = cat_tf.paragraphs[0]
        p_cat.text = data["category"]
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_GREEN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.733), Inches(0.8))
        title_tf = title_box.text_frame
        p_title = title_tf.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.85), Inches(11.733), Inches(0.5))
        sub_tf = sub_box.text_frame
        p_sub = sub_tf.paragraphs[0]
        p_sub.text = data["subtitle"]
        p_sub.font.size = Pt(14)
        p_sub.font.color.rgb = COLOR_MUTED

        # Content Card Box
        card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.5), Inches(11.733), Inches(4.2))
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = COLOR_CARD
        card_shape.line.color.rgb = COLOR_BORDER
        card_shape.line.width = Pt(1.5)

        # Points
        points_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(11.0), Inches(3.6))
        tf_points = points_box.text_frame
        tf_points.word_wrap = True

        for i, pt_text in enumerate(data["points"]):
            p = tf_points.paragraphs[0] if i == 0 else tf_points.add_paragraph()
            p.text = f"•  {pt_text}"
            p.font.size = Pt(16)
            p.font.color.rgb = COLOR_WHITE
            p.space_after = Pt(22)

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.3))
        ft_tf = footer_box.text_frame
        p_ft = ft_tf.paragraphs[0]
        p_ft.text = "CYBERFUSION XDR ENTERPRISE  |  Confidential Security Operations Platform  |  v3.4.0"
        p_ft.font.size = Pt(9)
        p_ft.font.color.rgb = RGBColor(100, 116, 139)

    output_path = "CYBERFUSION_XDR_ENTERPRISE.pptx"
    prs.save(output_path)
    print(f"PowerPoint Presentation successfully saved to: {output_path}")

if __name__ == "__main__":
    create_presentation()
