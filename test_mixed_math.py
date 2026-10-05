import docx
from docx.oxml import parse_xml
import latex2mathml.converter
import lxml.etree as ET

xslt_path = r'C:\Program Files\Microsoft Office\Office15\MML2OMML.XSL'
xslt = ET.parse(xslt_path)
transform = ET.XSLT(xslt)

def latex_to_omml_element(latex_str, is_inline=True):
    mathml = latex2mathml.converter.convert(latex_str)
    dom = ET.fromstring(mathml)
    new_dom = transform(dom)
    root = new_dom.getroot()
    xml_str = ET.tostring(root).decode('utf-8')
    if xml_str.startswith("<?xml"):
        xml_str = xml_str.split("?>", 1)[1].strip()
        
    if is_inline:
        return parse_xml(xml_str)
    else:
        wrapped = f'<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{xml_str}</m:oMathPara>'
        return parse_xml(wrapped)

doc = docx.Document()
p = doc.add_paragraph()
p.add_run("In this work, the decision threshold ")
p._p.append(latex_to_omml_element(r"\theta_h = P_{2.0}", is_inline=True))
p.add_run(" is calibrated on the baseline matrix ")
p._p.append(latex_to_omml_element(r"\mathbf{X}_h \in \mathbb{R}^{m \times 6}", is_inline=True))
p.add_run(" with risk formulation:")

# Block equation
p_block = doc.add_paragraph()
p_block.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
p_block._p.append(latex_to_omml_element(r"s(\mathbf{x}, n) = 2^{-\frac{\mathbb{E}[h(\mathbf{x})]}{c(n)}}", is_inline=False))

doc.save("test_mixed.docx")

import win32com.client
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
import os
d = word.Documents.Open(os.path.abspath("test_mixed.docx"))
d.SaveAs(os.path.abspath("test_mixed.pdf"), FileFormat=17)
d.Close()
word.Quit()
print("[OK] test_mixed.docx and test_mixed.pdf created and validated!")
