import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_faithful_docx(source_md_path, target_docx_path):
    print(f"Reading {source_md_path}...")
    with open(source_md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = docx.Document()

    # Set 0.5 inch margins to ensure wide 102-char ASCII diagrams NEVER wrap
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Set default style font to Consolas (exact monospace font preservation)
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Consolas'
    normal_style.font.size = Pt(8.5)
    normal_style.font.color.rgb = RGBColor(15, 23, 42)

    i = 0
    n = len(lines)
    in_code_block = False

    while i < n:
        raw_line = lines[i]
        line = raw_line.rstrip('\r\n')

        # Code block toggle (```)
        if line.strip().startswith("```"):
            in_code_block = not in_code_block
            # Add a subtle code fence marker line
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = Pt(10)
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.0)
            run.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        if in_code_block:
            # Code block / Diagram line: MUST preserve exact spaces, font, and 1.0 line spacing
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(9.5)  # exact line spacing to keep boxes tight
            
            run = p.add_run(line if line else " ")
            run.font.name = 'Consolas'
            run.font.size = Pt(7.8)  # 7.8pt Consolas ensures 102 characters fit in 7.5 inches without wrap
            
            # Accent colors for box borders vs text
            if any(char in line for char in ['┌', '┐', '└', '┘', '├', '┤', '┬', '┴', '┼', '─', '│', '+', '|', '=']):
                run.font.color.rgb = RGBColor(30, 58, 138)  # Indigo/Navy for crisp diagrams
                run.font.bold = False
            else:
                run.font.color.rgb = RGBColor(15, 23, 42)
            
            i += 1
            continue

        # Outside code block
        # Headings
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(2, 132, 199)
            i += 1
            continue

        if line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ["---", "***", "___"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("-" * 100)
            run.font.name = 'Consolas'
            run.font.size = Pt(7.5)
            run.font.color.rgb = RGBColor(203, 213, 225)
            i += 1
            continue

        # Standard lines (tables, lists, text)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = Pt(10.5)

        if line.strip().startswith("|") and line.strip().endswith("|"):
            # Table row in monospace
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(7.8)
            if re.match(r'^\|\s*:?-+:?\s*\|', line):
                run.font.color.rgb = RGBColor(148, 163, 184)
            elif "---" in line or "| **" in line:
                run.font.color.rgb = RGBColor(30, 58, 138)
                run.font.bold = True
            else:
                run.font.color.rgb = RGBColor(30, 41, 59)
        elif line.strip().startswith("* ") or line.strip().startswith("- "):
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif re.match(r'^\s*\d+\.\s+', line):
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif line.strip():
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
        else:
            # Blank line spacer
            p.paragraph_format.space_after = Pt(3)

        i += 1

    print(f"Saving faithful document to {target_docx_path}...")
    doc.save(target_docx_path)
    print(f"Successfully generated {target_docx_path} ({os.path.getsize(target_docx_path)} bytes)")

if __name__ == "__main__":
    src = "README_4.md"
    create_faithful_docx(src, "README_4.docx")
    create_faithful_docx(src, "ppt notes.docx")
