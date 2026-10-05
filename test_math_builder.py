import docx
from docx.oxml import parse_xml
import latex2mathml.converter
import lxml.etree as ET

xslt_path = r'C:\Program Files\Microsoft Office\Office15\MML2OMML.XSL'
xslt = ET.parse(xslt_path)
transform = ET.XSLT(xslt)

def latex_to_omml(latex_str, is_inline=False):
    mathml = latex2mathml.converter.convert(latex_str)
    dom = ET.fromstring(mathml)
    new_dom = transform(dom)
    root = new_dom.getroot()
    xml_str = ET.tostring(root).decode('utf-8')
    
    # Strip any <?xml...?> declaration
    if xml_str.startswith("<?xml"):
        xml_str = xml_str.split("?>", 1)[1].strip()
        
    if is_inline:
        return parse_xml(xml_str)
    else:
        wrapped = f'<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{xml_str}</m:oMathPara>'
        return parse_xml(wrapped)

doc = docx.Document()
p = doc.add_paragraph("Testing Block Equation:")

latex_eq = r'\theta_h = P_{2.0}\left( \{\text{Score}_{\text{iso}}(\mathbf{x}_i)\}_{i=1}^{50} \right)'
elem = latex_to_omml(latex_eq, is_inline=False)
p._p.append(elem)

p2 = doc.add_paragraph("Testing Inline Equation: ")
p2.add_run("The threshold is ")
inline_elem = latex_to_omml(r'\theta_h', is_inline=True)
p2._p.append(inline_elem)
p2.add_run(" and risk is ")
inline_elem2 = latex_to_omml(r'R(\mathbf{x}) \in [0, 100]', is_inline=True)
p2._p.append(inline_elem2)

doc.save("test_math.docx")
print("[OK] test_math.docx successfully generated with native OMML equations!")

# Now test PDF conversion via Word COM
import win32com.client
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False
import os
doc_word = word.Documents.Open(os.path.abspath("test_math.docx"))
doc_word.SaveAs(os.path.abspath("test_math.pdf"), FileFormat=17)
doc_word.Close()
word.Quit()
print("[OK] test_math.pdf successfully generated via Word!")
