import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_formatted_text(paragraph, text, default_color=RGBColor(30, 41, 59), font_size=Pt(10.5)):
    # Regular expression to tokenize bold, italic, inline code, and links
    # Pattern handles: [link](url), **bold**, *italic*, `code`
    token_pattern = re.compile(
        r'(\[.*?\]\(.*?\)|\*\*.*?\*\*|\*.*?\*|`.*?`)'
    )
    parts = token_pattern.split(text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = "Segoe UI"
        run.font.size = font_size
        run.font.color.rgb = default_color
        
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            run.text = part[1:-1]
            run.italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run.text = part[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 40, 70)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            m = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if m:
                run.text = m.group(1)
                run.font.color.rgb = RGBColor(2, 132, 199)
                run.underline = True
            else:
                run.text = part
        else:
            run.text = part

def build_docx_from_readme(readme_path, output_docx_path):
    print(f"Reading from {readme_path}...")
    with open(readme_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = docx.Document()

    # Page Margins: 1 inch (72 pt = 1440 dxa)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set base Normal style font
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Segoe UI'
    font_normal.size = Pt(10.5)
    font_normal.color.rgb = RGBColor(30, 41, 59)

    i = 0
    n = len(lines)
    in_code_block = False
    code_block_lines = []

    while i < n:
        line = lines[i].rstrip('\r\n')
        
        # Check code block toggle
        if line.strip().startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_block_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                # Emit Code Block / Diagram
                full_code_text = "\n".join(code_block_lines)
                
                # Create a single cell table with light background for code/diagram
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = False
                
                cell = tbl.cell(0, 0)
                cell.width = Inches(6.5)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                
                # Set light border around code block
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
                    f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0284C7"/>'
                    f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
                    f'<w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
                    f'</w:tcBorders>'
                )
                tcPr.append(tcBorders)

                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                cp.paragraph_format.line_spacing = 1.05
                
                c_run = cp.add_run(full_code_text)
                c_run.font.name = "Consolas"
                c_run.font.size = Pt(8.0)
                c_run.font.color.rgb = RGBColor(15, 23, 42)
                
                # Spacer paragraph after code block
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(0)
                p_spacer.paragraph_format.space_after = Pt(4)
                
                code_block_lines = []
                i += 1
                continue
        
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Check Table
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table lines
            rows_data = []
            for tl in table_lines:
                # Filter out separator lines like | :--- | :--- |
                inner = tl[1:-1]
                cols = [c.strip() for c in inner.split("|")]
                if all(re.match(r'^:?-+:?$', c) for c in cols if c):
                    continue
                rows_data.append(cols)
            
            if rows_data:
                num_cols = max(len(r) for r in rows_data)
                # Pad shorter rows
                for r in rows_data:
                    while len(r) < num_cols:
                        r.append("")
                
                table = doc.add_table(rows=len(rows_data), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(table, color="CBD5E1", sz="4")
                
                for row_idx, r_data in enumerate(rows_data):
                    is_header = (row_idx == 0)
                    row = table.rows[row_idx]
                    
                    # Set cantSplit on row
                    trPr = row._tr.get_or_add_trPr()
                    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                    if is_header:
                        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

                    for col_idx, text in enumerate(r_data):
                        cell = row.cells[col_idx]
                        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                        
                        if is_header:
                            set_cell_background(cell, "1E3A8A") # Navy Blue
                        else:
                            if row_idx % 2 == 1:
                                set_cell_background(cell, "FFFFFF")
                            else:
                                set_cell_background(cell, "F8FAFC") # Soft alternate slate
                        
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.1
                        
                        cell_color = RGBColor(255, 255, 255) if is_header else RGBColor(30, 41, 59)
                        f_size = Pt(9.5) if not is_header else Pt(9.5)
                        
                        if is_header:
                            run = p.add_run(text)
                            run.font.name = "Segoe UI"
                            run.font.size = f_size
                            run.font.color.rgb = cell_color
                            run.bold = True
                        else:
                            add_formatted_text(p, text, default_color=cell_color, font_size=f_size)
                
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(0)
                p_spacer.paragraph_format.space_after = Pt(6)
            continue

        # Check Horizontal Rule
        if line.strip() in ["---", "***", "___"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pbdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            pPr.append(pbdr)
            i += 1
            continue

        # Check Headings
        if line.startswith("# "):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(title_text)
            run.font.name = "Segoe UI"
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Slate-900
            i += 1
            continue

        if line.startswith("## "):
            h1_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h1_text)
            run.font.name = "Segoe UI"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Blue-900
            
            # Subtle accent underline
            pPr = p._p.get_or_add_pPr()
            pbdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="0284C7"/></w:pBdr>')
            pPr.append(pbdr)
            
            i += 1
            continue

        if line.startswith("### "):
            h2_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h2_text)
            run.font.name = "Segoe UI"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(2, 132, 199) # Sky-600
            i += 1
            continue

        if line.startswith("#### "):
            h3_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h3_text)
            run.font.name = "Segoe UI"
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        # Check Bullet Lists (* or -)
        if line.strip().startswith("* ") or line.strip().startswith("- "):
            list_text = re.sub(r'^\s*[\*\-]\s+', '', line)
            indent_level = len(line) - len(line.lstrip())
            
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if indent_level > 2:
                p.paragraph_format.left_indent = Inches(0.5)
            else:
                p.paragraph_format.left_indent = Inches(0.25)
            
            add_formatted_text(p, list_text, default_color=RGBColor(30, 41, 59))
            i += 1
            continue

        # Check Numbered Lists (e.g. 1. , 2. )
        m_num = re.match(r'^\s*(\d+)\.\s+(.*)$', line)
        if m_num:
            num_val = m_num.group(1)
            num_text = m_num.group(2)
            
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.line_spacing = 1.15
            
            run_num = p.add_run(f"{num_val}. ")
            run_num.font.name = "Segoe UI"
            run_num.font.bold = True
            run_num.font.color.rgb = RGBColor(30, 58, 138)
            
            add_formatted_text(p, num_text, default_color=RGBColor(30, 41, 59))
            i += 1
            continue

        # Regular Paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, line.strip(), default_color=RGBColor(30, 41, 59))
        else:
            # Empty line spacer
            pass

        i += 1

    print(f"Saving to {output_docx_path}...")
    doc.save(output_docx_path)
    print(f"Successfully generated {output_docx_path} ({os.path.getsize(output_docx_path)} bytes)")

if __name__ == "__main__":
    src = "README_4.md"
    dst = "ppt notes.docx"
    build_docx_from_readme(src, dst)
